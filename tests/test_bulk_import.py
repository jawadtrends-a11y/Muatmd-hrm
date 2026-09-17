"""
حرّاس الاستيراد الجماعيّ (ق-154).

**شركةٌ بمئة موظف لا يُعقل إدخالهم واحدًا واحدًا** (قرار جواد).

⚠️⚠️ **والخطأ في سطرٍ يوقف الملفّ كلّه**: فاستيرادٌ نصفُه ناقصٌ
**أسوأ من لا شيء** — والشركة لا تعرف من دخل ومن بقي.
"""
import io
from datetime import date

import pytest

from apps.accounts.models import Account, Company
from apps.accounts.services.provisioning import provision_account
from apps.core.tenancy.context import account_scope
from apps.employees.models import Employment, Person
from apps.employees.services import bulk_import as imp

HEAD = ("الرقم الوظيفي,الاسم الأول,اسم العائلة,الجنس,الجنسية,"
        "نوع الهوية,رقم الهوية,تاريخ المباشرة,الجوال,البريد,"
        "الإدارة,المسمى الوظيفي,الراتب الأساسي,بدل السكن,"
        "بدل النقل,الآيبان")

OK1 = ("5001,طلال,السبيعي,male,SA,national_id,1066554433,"
       "2026-01-15,0501112233,,,,9000,2000,800,")
OK2 = ("5002,ريما,الدوسري,female,SA,national_id,1055443322,"
       "2026-02-01,0502223344,,,,8500,1500,700,")


@pytest.fixture
def env(db):
    r = provision_account(slug="imp-t", display_name_ar="حساب",
                          company_name_ar="شركة", is_sandbox=True)
    with account_scope(r.account_id):
        yield {"account_id": r.account_id,
               "comp": Company.objects.get(id=r.company_id)}


def _parse(env, rows):
    content = (HEAD + "\n" + "\n".join(rows)).encode("utf-8-sig")
    return imp.parse_file(content, env["comp"])


def test_template_has_all_columns(env):
    """والقالب يحوي كل الأعمدة بعناوينها العربية."""
    text = imp.template_csv()
    head = text.split("\n")[0]
    for _, label, required, _x in imp.COLUMNS:
        if required:
            assert label in head, label


def test_valid_file_parses(env):
    """والملفّ الصحيح يُقرأ ويُعدّ."""
    with account_scope(env["account_id"]):
        out = _parse(env, [OK1, OK2])
        assert out["valid"] == 2
        assert out["invalid"] == 0
        assert out["can_import"] is True


def test_one_error_blocks_the_whole_file(env):
    """
    ⚠️⚠️ الأهمّ: **خطأٌ واحد يوقف الملفّ كلّه** (قرار جواد).

    فاستيرادٌ نصفُه ناقصٌ أسوأ من لا شيء.
    """
    with account_scope(env["account_id"]):
        bad = ("5003,سعود,الحربي,ذكر,SA,national_id,1044332211,"
               "2026-01-15,,,,,7000,,,")
        out = _parse(env, [OK1, OK2, bad])
        assert out["valid"] == 2
        assert out["invalid"] == 1
        assert out["can_import"] is False, (
            "سُمح بالاستيراد وفي الملفّ خطأ")


def test_all_errors_are_returned(env):
    """
    ⚠️ **وتُرجَع كل الأخطاء لا أوّلها**: فمن يُصلح خطأً ثم يرفع
    ليكتشف ثانيًا **يكرّر الرفع عشرًا**.
    """
    with account_scope(env["account_id"]):
        rows = [
            "5004,ن,ا,ذكر,SA,national_id,1033221100,2026-01-15,,,,,,,,",
            "5005,م,ب,female,SA,national_id,1022110099,99-99-9999,,,,,,,,",
        ]
        out = _parse(env, rows)
        assert out["invalid"] == 2, out
        assert len(out["errors"]) == 2


def test_duplicate_in_file_is_caught(env):
    """والتكرار داخل الملفّ يُكشف — قبل أن يصل القاعدة."""
    with account_scope(env["account_id"]):
        out = _parse(env, [OK1, OK1])
        assert out["invalid"] >= 1
        joined = " ".join(e["errors"][0] for e in out["errors"])
        assert "مكرّر" in joined


def test_duplicate_with_existing_is_caught(env):
    """والمسجَّل في النظام يُكشف كذلك."""
    with account_scope(env["account_id"]):
        imp.execute(company=env["comp"],
                    parsed_rows=_parse(env, [OK1])["rows"])
        out = _parse(env, [OK1])
        assert out["can_import"] is False
        joined = " ".join(e["errors"][0] for e in out["errors"])
        assert "مستعمل" in joined or "مسجَّل" in joined


def test_unknown_department_is_created(env):
    """
    ق-215: ⚠️⚠️ **وإدارةٌ غير معرَّفة تُنشأ** (قرار جواد) — ولا
    تُعدّ خطأً.

    ⚠️ **والمعاينة تُنبّه بما سيُنشأ**: فالخطأ الإملائيّ **يصير
    إدارةً جديدة**، والعميل يراجع القائمة قبل التنفيذ.
    """
    with account_scope(env["account_id"]):
        row = ("5006,فهد,العمري,male,SA,national_id,1011009988,"
               "2026-01-15,,,إدارةٌ جديدة,مسمّى جديد,7000,,,")
        out = _parse(env, [row])
        assert out["invalid"] == 0, out["errors"]
        assert "إدارةٌ جديدة" in out["new_departments"]
        assert "مسمّى جديد" in out["new_job_titles"]


def test_new_reference_is_created_once(env):
    """
    ⚠️⚠️ الأهمّ: **وبلا تكرار** (قرار جواد): **فإدارةٌ على ثلاثة
    صفوفٍ تُنشأ مرّةً واحدة**.
    """
    from apps.employees.services.bulk_import import execute
    from apps.organization.models import Department, JobTitle

    with account_scope(env["account_id"]):
        rows = [
            (f"520{i},موظف{i},التجربة,male,SA,national_id,"
             f"10770088{i}{i},2026-01-15,,,إدارة مشتركة,"
             f"مسمّى مشترك,7000,,,")
            for i in (1, 2, 3)
        ]
        out = _parse(env, rows)
        assert out["invalid"] == 0, out["errors"]
        # ⚠️ **وتُعرَض مرّةً لا ثلاثًا**
        assert out["new_departments"].count("إدارة مشتركة") == 1

        execute(company=env["comp"], parsed_rows=out["rows"])

        assert Department.objects.filter(
            company=env["comp"], name_ar="إدارة مشتركة").count() == 1
        assert JobTitle.objects.filter(
            company=env["comp"], name_ar="مسمّى مشترك").count() == 1


def test_execute_creates_employees(env):
    """والتنفيذ يُنشئ الموظفين بهيكل رواتبهم."""
    with account_scope(env["account_id"]):
        out = _parse(env, [OK1, OK2])
        res = imp.execute(company=env["comp"],
                          parsed_rows=out["rows"])
        assert res["created"] == 2

        e = Employment.objects.filter(company=env["comp"],
                                      employee_no="5001").first()
        assert e is not None
        assert e.join_date == date(2026, 1, 15)
        assert e.person.family_name_ar == "السبيعي"


def test_failure_rolls_everything_back(env):
    """
    ⚠️⚠️ **وإخفاق سطرٍ يُرجع الجميع**: وإلا بقي نصفُ الملفّ ولا
    يُعرف أين توقّف.
    """
    with account_scope(env["account_id"]):
        before = Employment.objects.filter(
            company=env["comp"]).count()

        out = _parse(env, [OK1, OK2])
        # ⚠️ نُفسد الثاني بعد الفحص — كما لو تغيّر بين المعاينة
        # والتنفيذ
        out["rows"][1]["employee_no"] = ""

        with pytest.raises(imp.ImportError_):
            imp.execute(company=env["comp"],
                        parsed_rows=out["rows"])

        after = Employment.objects.filter(
            company=env["comp"]).count()
        assert after == before, "بقي نصفُ الملفّ في النظام"


def test_empty_file_refused(env):
    """وملفٌّ فارغ يُرفض برسالةٍ مفهومة."""
    with account_scope(env["account_id"]):
        with pytest.raises(imp.ImportError_) as e:
            imp.parse_file(HEAD.encode("utf-8-sig"), env["comp"])
        assert "فارغ" in str(e.value)


def test_missing_required_column_refused(env):
    """وعمودٌ إلزاميّ ناقص يُرفض قبل قراءة الصفوف."""
    with account_scope(env["account_id"]):
        head = "الرقم الوظيفي,الاسم الأول"
        with pytest.raises(imp.ImportError_) as e:
            imp.parse_file((head + "\n1,س").encode("utf-8-sig"),
                           env["comp"])
        assert "ناقصة" in str(e.value)


def test_example_row_is_skipped(env):
    """
    وسطر المثال في القالب يُتخطّى بلا خطأ — فمن نزّل القالب
    ورفعه فارغًا لا يرى خطأً وهميًّا.
    """
    with account_scope(env["account_id"]):
        out = imp.parse_file(
            imp.template_csv().encode("utf-8-sig"), env["comp"])
        assert out["total"] == 0
        assert out["errors"] == []


def test_template_has_dropdowns(env):
    """
    ⚠️ **وقوائمُ منسدلة للحقول الملتبسة** (بلاغ جواد).

    فبلا قائمةٍ يكتب العميل «ذكر» فيُرفض الملفّ، **ويُعيد الرفع
    مرارًا**.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(imp.template_xlsx()))
    ws = wb.active
    dvs = list(ws.data_validations.dataValidation)
    assert len(dvs) >= 3, f"قوائم ناقصة: {len(dvs)}"

    formulas = " ".join((d.formula1 or "") for d in dvs)
    assert "male" in formulas and "female" in formulas
    assert "national_id" in formulas and "iqama" in formulas

    # ⚠️ والجنسيات في ورقةٍ مرجعية مخفيّة
    assert "مرجع" in wb.sheetnames
    assert wb["مرجع"].sheet_state == "hidden", (
        "ورقة المرجع ظاهرة — تُربك من يملأ")


def test_nationality_is_not_restricted(env):
    """
    ⚠️ **ولا تُمنع جنسيةٌ خارج القائمة**: فمن لم يجد جنسيّته
    يكتبها — والقائمة إرشادٌ لا سجن.
    """
    with account_scope(env["account_id"]):
        row = ("7301,أنطون,بيريز,male,MX,iqama,2077003300,"
               "2026-03-01,,,,,7000,,,")
        out = _parse(env, [row])
        assert out["valid"] == 1, out["errors"]

