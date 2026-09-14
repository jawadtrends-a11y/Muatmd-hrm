"""
حرّاس استيراد أرصدة الإجازات (ق-170).

**بلاغ جواد:** شركةٌ تنتقل من نظامٍ آخر **لا تُعيد إدخال أرصدة
موظفيها واحدًا واحدًا**.

⚠️⚠️ **والرصيد افتتاحيٌّ بتاريخه** — والاستحقاق اليوميّ **يُضاف
إليه**، **لا يحلّ محلّه**.
"""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook

from apps.accounts.models import Account, Company
from apps.accounts.services.provisioning import provision_account
from apps.core.tenancy.context import account_scope
from apps.employees.services.hiring import create_employment, create_person
from apps.leaves.models import LeaveBalance, LeaveType
from apps.leaves.services import import_balances as imp
from apps.payroll.models import PayComponent

HEAD = ["الرقم الوظيفي", "الرصيد المستحق", "الرصيد كما في"]


@pytest.fixture
def env(db):
    r = provision_account(slug="bal-t", display_name_ar="حساب",
                          company_name_ar="شركة", is_sandbox=True)
    with account_scope(r.account_id):
        acc = Account.objects.get(id=r.account_id)
        comp = Company.objects.get(id=r.company_id)
        basic = PayComponent.objects.get(company=comp, code="BASIC")

        for i, (fn, ln, nid, mob) in enumerate((
                ("راكان", "العصيمي", "1077001100", "0507700110"),
                ("جواهر", "الثبيتي", "1077002200", "0507700220"))):
            p, _ = create_person(
                account=acc, first_name_ar=fn, family_name_ar=ln,
                gender="male", nationality_code="SA",
                id_type="national_id", id_number=nid, mobile=mob)
            create_employment(
                person=p, company=comp, employee_no=f"B-{i}",
                join_date=date(2024, 1, 1),
                salary_lines=[(basic, Decimal("9000"))])

        lt = LeaveType.objects.filter(company=comp,
                                      code="ANNUAL").first()
        yield {"account_id": r.account_id, "comp": comp, "lt": lt}


def _file(rows):
    wb = Workbook()
    ws = wb.active
    for i, h in enumerate(HEAD, 1):
        ws.cell(1, i, h)
    for j, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            ws.cell(j, i, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(env, rows):
    return imp.parse_file(_file(rows), env["comp"], env["lt"])


def test_template_is_xlsx(env):
    """
    ⚠️ **والقالب xlsx لا CSV** — فإكسل العربيّ يفتح CSV بالفواصل
    **في عمودٍ واحد** (درس ق-١٥٥).
    """
    data = imp.template_xlsx()
    assert data[:2] == b"PK"


def test_empty_template_parses_clean(env):
    """
    ⚠️⚠️ **وسطر المثال يُتخطّى** — فمن نزّل القالب ورفعه فارغًا
    **لا يرى خطأً وهميًّا**.
    """
    with account_scope(env["account_id"]):
        out = imp.parse_file(imp.template_xlsx(), env["comp"],
                             env["lt"])
        assert out["total"] == 0
        assert out["errors"] == []


def test_real_row_matching_example_is_not_dropped(env):
    """
    ⚠️⚠️ **وصفٌّ حقيقيٌّ يطابق قيم المثال لا يُحذف**: فمطابقةُ
    القيم **تحذف صفًّا صامتًا** — والعلامة تحسمها.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["B-0", "12.5", "2025-12-31"]])
        assert out["total"] == 1, "حُذف صفٌّ حقيقيّ"
        assert out["valid"] == 1


def test_unknown_employee_is_an_error(env):
    """ورقمٌ غير موجود خطأ."""
    with account_scope(env["account_id"]):
        out = _parse(env, [["ZZZ", "10", "2025-12-31"]])
        assert out["invalid"] == 1
        assert "غير موجود" in out["errors"][0]["errors"][0]


def test_negative_balance_refused(env):
    """
    ⚠️ **ورصيدٌ سالب يُرفض**: فمن استهلك أكثر ممّا استحقّ
    **يُسوّى قبل النقل لا بعده**.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["B-0", "-4", "2025-12-31"]])
        assert out["invalid"] == 1
        assert "سالب" in out["errors"][0]["errors"][0]


def test_duplicate_in_file_is_caught(env):
    """والتكرار داخل الملفّ يُكشف."""
    with account_scope(env["account_id"]):
        out = _parse(env, [["B-0", "10", "2025-12-31"],
                           ["B-0", "12", "2025-12-31"]])
        assert out["invalid"] >= 1


def test_one_error_blocks_the_file(env):
    """
    ⚠️⚠️ **وخطأٌ واحد يوقف الملفّ كلّه** (قرار جواد): فاستيرادٌ
    نصفُه ناقصٌ **أسوأ من لا شيء**.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["B-0", "10", "2025-12-31"],
                           ["ZZZ", "5", "2025-12-31"]])
        assert out["valid"] == 1
        assert out["can_import"] is False


def test_execute_writes_opening_balance(env):
    """
    ⚠️⚠️ الأهمّ: **والرصيد يُكتب افتتاحيًّا لا مستحقًّا** — فهو
    ما جاء من النظام السابق.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["B-0", "12.5", "2025-12-31"],
                           ["B-1", "8", "2025-12-31"]])
        res = imp.execute(company=env["comp"], leave_type=env["lt"],
                          parsed_rows=out["rows"])
        assert res["imported"] == 2

        bal = LeaveBalance.objects.filter(
            employment__employee_no="B-0", year=2025).first()
        assert bal is not None
        assert bal.opening_balance == Decimal("12.5")
        # ⚠️ **والمستحقّ يبقى صفرًا** — فالاستحقاق يُضاف لاحقًا
        assert bal.accrued == Decimal("0")


def test_accrual_date_starts_from_as_of(env):
    """
    ⚠️⚠️ **والاحتساب يبدأ من تاريخه**: فلا يُعاد استحقاقُ ما
    قبله — **وإلا ضوعف الرصيد**.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["B-0", "12.5", "2025-12-31"]])
        imp.execute(company=env["comp"], leave_type=env["lt"],
                    parsed_rows=out["rows"])
        bal = LeaveBalance.objects.filter(
            employment__employee_no="B-0", year=2025).first()
        assert bal.last_accrual_date == date(2025, 12, 31)


def test_reimport_updates_not_duplicates(env):
    """
    **وإعادة الاستيراد تُحدّث ولا تُكرّر** — فملفٌّ يُرفع مرّتين
    **لا يُضاعف الرصيد**.
    """
    with account_scope(env["account_id"]):
        for val in ("10", "15"):
            out = _parse(env, [["B-0", val, "2025-12-31"]])
            imp.execute(company=env["comp"], leave_type=env["lt"],
                        parsed_rows=out["rows"])

        rows = LeaveBalance.objects.filter(
            employment__employee_no="B-0", year=2025)
        assert rows.count() == 1, "تكرّر الرصيد"
        assert rows.first().opening_balance == Decimal("15")


def test_semicolon_csv_is_parsed(env):
    """وCSV بالفاصلة المنقوطة يُقرأ — فمن حفظ من إكسل العربيّ."""
    with account_scope(env["account_id"]):
        text = (";".join(HEAD) + "\n"
                + "B-0;9;2025-12-31")
        out = imp.parse_file(text.encode("utf-8-sig"), env["comp"],
                             env["lt"])
        assert out["valid"] == 1, out["errors"]
