"""
حرّاس استيراد الحضور (ق-171).

**بلاغ جواد:** شركةٌ تنتقل من نظامٍ آخر **لا تُعيد إدخال سجلّ
حضور موظفيها يومًا يومًا**.

⚠️⚠️ **ولا يُستورَد يومٌ في مسيرٍ معتمد**: فالأجر احتُسب عليه،
**وتغييرُه يجعل القسيمة لا تفسّر نفسها**.
"""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook

from apps.accounts.models import Account, Company
from apps.accounts.services.provisioning import provision_account
from apps.attendance.models import AttendanceDay
from apps.attendance.services import import_days as imp
from apps.core.tenancy.context import account_scope
from apps.employees.services.hiring import create_employment, create_person
from apps.payroll.models import (
    PayComponent, PayrollRunStatus, PayrollRunType)
from apps.payroll.services import engine

HEAD = ["الرقم الوظيفي", "التاريخ", "الحالة", "الدخول", "الخروج"]


@pytest.fixture
def env(db):
    r = provision_account(slug="day-t", display_name_ar="حساب",
                          company_name_ar="شركة", is_sandbox=True)
    with account_scope(r.account_id):
        acc = Account.objects.get(id=r.account_id)
        comp = Company.objects.get(id=r.company_id)
        basic = PayComponent.objects.get(company=comp, code="BASIC")

        for i, (fn, ln, nid, mob) in enumerate((
                ("صالح", "الحميدي", "1078001100", "0507800110"),
                ("دلال", "الرشود", "1078002200", "0507800220"))):
            p, _ = create_person(
                account=acc, first_name_ar=fn, family_name_ar=ln,
                gender="male", nationality_code="SA",
                id_type="national_id", id_number=nid, mobile=mob)
            create_employment(
                person=p, company=comp, employee_no=f"D-{i}",
                join_date=date(2024, 1, 1),
                salary_lines=[(basic, Decimal("9000"))])

        yield {"account_id": r.account_id, "comp": comp}


def _file(rows):
    wb = Workbook()
    ws = wb.active
    for i, h in enumerate(HEAD, 1):
        ws.cell(1, i, h)
    for j, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            ws.cell(j, i, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(env, rows):
    return imp.parse_file(_file(rows), env["comp"])


def test_template_is_xlsx_with_dropdown(env):
    """
    ⚠️ **وقائمةٌ منسدلة للحالة** — فالعميل قد لا يعرف الصيغة
    (درس ق-١٥٥).
    """
    from openpyxl import load_workbook

    data = imp.template_xlsx()
    assert data[:2] == b"PK"
    ws = load_workbook(io.BytesIO(data)).active
    dvs = list(ws.data_validations.dataValidation)
    assert dvs, "لا قائمة منسدلة"
    assert "present" in (dvs[0].formula1 or "")


def test_valid_rows_parse(env):
    """والصحيح يُقرأ."""
    with account_scope(env["account_id"]):
        out = _parse(env, [
            ["D-0", "2025-11-03", "present", "08:00", "17:00"],
            ["D-1", "2025-11-03", "absent", "", ""]])
        assert out["valid"] == 2, out["errors"]
        assert out["can_import"] is True


def test_arabic_status_is_accepted(env):
    """**والحالة بالعربية تُقبل** — فمن يكتب «حاضر» يُفهم."""
    with account_scope(env["account_id"]):
        out = _parse(env, [["D-0", "2025-11-03", "حاضر", "", ""]])
        assert out["valid"] == 1, out["errors"]
        assert out["rows"][0]["status"] == "present"


def test_unknown_status_is_an_error(env):
    """
    ⚠️ **وحالةٌ مخترَعة خطأ**: فهي **تُفسد التقارير صامتة**.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["D-0", "2025-11-03", "ذكر", "", ""]])
        assert out["invalid"] == 1
        assert "حالة غير معروفة" in out["errors"][0]["errors"][0]


def test_checkout_before_checkin_refused(env):
    """
    ⚠️ **وخروجٌ قبل دخول خطأ** — فاليوم لا يُقرأ عكسًا.
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["D-0", "2025-11-03", "present",
                            "17:00", "08:00"]])
        assert out["invalid"] == 1
        assert "الخروج قبل الدخول" in out["errors"][0]["errors"][0]


def test_duplicate_day_is_caught(env):
    """ويومٌ مكرّرٌ لموظفٍ يُكشف — فلا يُكتب مرّتين."""
    with account_scope(env["account_id"]):
        out = _parse(env, [["D-0", "2025-11-03", "present", "", ""],
                           ["D-0", "2025-11-03", "absent", "", ""]])
        assert out["invalid"] >= 1


def test_approved_payroll_month_is_locked(env):
    """
    ⚠️⚠️ الأهمّ: **ولا يُستورَد يومٌ في مسيرٍ معتمد** — فالأجر
    احتُسب عليه، **وتغييرُه يجعل القسيمة لا تفسّر نفسها**.
    """
    with account_scope(env["account_id"]):
        run = engine.create_run(company=env["comp"],
                                run_type=PayrollRunType.REGULAR,
                                year=2025, month=11)
        engine.calculate_run(run)
        run.status = PayrollRunStatus.SUBMITTED
        run.save(update_fields=["status"])
        engine.approve_run(run, None)

        out = _parse(env, [["D-0", "2025-11-03", "present", "", ""]])
        assert out["invalid"] == 1, out
        assert "مسيرٍ معتمد" in out["errors"][0]["errors"][0]


def test_open_month_is_allowed(env):
    """**وشهرٌ بلا مسيرٍ معتمد يُستورَد**."""
    with account_scope(env["account_id"]):
        out = _parse(env, [["D-0", "2025-10-03", "present", "", ""]])
        assert out["valid"] == 1, out["errors"]


def test_one_error_blocks_the_file(env):
    """
    ⚠️⚠️ **وخطأٌ واحد يوقف الملفّ كلّه** (قرار جواد).
    """
    with account_scope(env["account_id"]):
        out = _parse(env, [["D-0", "2025-10-03", "present", "", ""],
                           ["ZZZ", "2025-10-03", "present", "", ""]])
        assert out["valid"] == 1
        assert out["can_import"] is False


def test_execute_writes_days(env):
    """والتنفيذ يكتب الأيام بأوقاتها."""
    with account_scope(env["account_id"]):
        out = _parse(env, [
            ["D-0", "2025-10-03", "present", "08:00", "17:00"]])
        res = imp.execute(company=env["comp"],
                          parsed_rows=out["rows"])
        assert res["imported"] == 1

        d = AttendanceDay.objects.filter(
            employment__employee_no="D-0",
            work_date=date(2025, 10, 3)).first()
        assert d is not None
        assert d.status == "present"
        # ⚠️ **والحقل تاريخٌ ووقت** — فيُبنى من اليوم والساعة.
        #
        # ⚠️⚠️ **ويُخزَّن بـUTC**: فالمقارنة بالساعة الخام تقرأ
        # ٥ بدل ٨ — **والتحويل للتوقيت المحلّيّ هو الصواب**.
        from django.utils import timezone

        assert d.first_in is not None
        local = timezone.localtime(d.first_in)
        assert local.hour == 8, local


def test_reimport_updates_not_duplicates(env):
    """
    **وإعادة الاستيراد تُحدّث ولا تُكرّر** — فملفٌّ يُرفع مرّتين
    **لا يُنشئ يومين**.
    """
    with account_scope(env["account_id"]):
        for st in ("present", "absent"):
            out = _parse(env, [["D-0", "2025-10-04", st, "", ""]])
            imp.execute(company=env["comp"], parsed_rows=out["rows"])

        rows = AttendanceDay.objects.filter(
            employment__employee_no="D-0",
            work_date=date(2025, 10, 4))
        assert rows.count() == 1, "تكرّر اليوم"
        assert rows.first().status == "absent"


def test_empty_template_parses_clean(env):
    """وسطر المثال يُتخطّى — فالقالب الفارغ بلا أخطاء."""
    with account_scope(env["account_id"]):
        out = imp.parse_file(imp.template_xlsx(), env["comp"])
        assert out["total"] == 0
        assert out["errors"] == []
