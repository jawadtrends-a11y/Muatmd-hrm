"""
استيراد أرصدة الإجازات (ق-170).

**بلاغ جواد:** شركةٌ تنتقل من نظامٍ آخر **لا تُعيد إدخال أرصدة
موظفيها واحدًا واحدًا**.

⚠️⚠️ **والرصيد افتتاحيٌّ بتاريخه** — والاستحقاق اليوميّ **يُضاف
إليه** من ذلك التاريخ، **لا يحلّ محلّه**.

⚠️ **والخطأ في سطرٍ يوقف الملفّ كلّه** (قرار جواد): فاستيرادٌ
نصفُه ناقصٌ **أسوأ من لا شيء**.
"""
import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction


class BalanceImportError(Exception):
    """سببٌ يُعرض للمستخدم كما هو."""


COLUMNS = [
    ("employee_no", "الرقم الوظيفي", True, "مثال: 1001"),
    ("balance", "الرصيد المستحق", True, "12.5"),
    ("as_of", "الرصيد كما في", True, "2025-12-31"),
]


def template_xlsx(company=None):
    """
    قالبٌ بعناوينه ومثالٍ يُحذف.

    ⚠️ **وxlsx لا CSV**: فإكسل العربيّ يفتح CSV بالفواصل **في
    عمودٍ واحد** (درس ق-١٥٥).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "أرصدة الإجازات"
    ws.sheet_view.rightToLeft = True

    head = PatternFill("solid", start_color="1F4E5F")
    req = PatternFill("solid", start_color="C0392B")

    for i, (key, label, required, example) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = req if required else head
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(
            len(label) + 8, 20)
        ex = ws.cell(row=2, column=i, value=example)
        ex.font = Font(italic=True, color="888888", size=10)
        ex.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_to_csv(data):
    """يحوّل xlsx لنصٍّ يُقرأ بالمنطق نفسه."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    buf = io.StringIO()
    w = csv.writer(buf)
    for row in ws.iter_rows(values_only=True):
        cells = []
        for v in row:
            if v is None:
                cells.append("")
            elif isinstance(v, datetime):
                cells.append(v.date().isoformat())
            elif isinstance(v, date):
                cells.append(v.isoformat())
            elif isinstance(v, float) and v.is_integer():
                cells.append(str(int(v)))
            else:
                cells.append(str(v).strip())
        w.writerow(cells)
    return buf.getvalue()


def parse_file(content, company, leave_type):
    """
    يقرأ الملفّ ويفحصه سطرًا سطرًا — **بلا كتابة**.

    ⚠️ **ويُرجع كل الأخطاء لا أوّلها**: فمن يُصلح خطأً ثم يرفع
    ليكتشف ثانيًا **يكرّر الرفع عشرًا**.
    """
    from apps.employees.models import Employment, EmploymentStatus

    if isinstance(content, bytes) and content[:2] == b"PK":
        content = _xlsx_to_csv(content)
    if isinstance(content, bytes):
        for enc in ("utf-8-sig", "utf-8", "cp1256"):
            try:
                content = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise BalanceImportError(
                "تعذّر قراءة الملفّ — احفظه بترميز UTF-8")

    first = content.split("\n", 1)[0]
    delim = ";" if first.count(";") > first.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(content), delimiter=delim)
            if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise BalanceImportError("الملفّ فارغ")

    headers = [h.strip() for h in rows[0]]
    by_label = {c[1]: c[0] for c in COLUMNS}
    mapping = {by_label[h]: i for i, h in enumerate(headers)
               if h in by_label}

    missing = [c[1] for c in COLUMNS if c[2] and c[0] not in mapping]
    if missing:
        raise BalanceImportError(
            "أعمدة إلزامية ناقصة: " + "، ".join(missing))

    emps = {e.employee_no: e for e in Employment.objects.filter(
        company=company, status=EmploymentStatus.ACTIVE)}

    parsed, errors, seen = [], [], set()
    for n, raw in enumerate(rows[1:], start=2):
        def cell(k):
            i = mapping.get(k)
            return (raw[i].strip() if i is not None and i < len(raw)
                    else "")

        no = cell("employee_no")
        # ⚠️⚠️ **وسطر المثال يُعلَّم بعلامةٍ لا تلتبس ببيانات**:
        # فمطابقةُ القيم **تحذف صفًّا حقيقيًّا صامتًا** — ومن رقمه
        # ١٠٠١ ورصيده ١٢٫٥ يضيع.
        if no.startswith("مثال"):
            continue

        row_errors = []
        if not no:
            row_errors.append("«الرقم الوظيفي» مطلوب")
        elif no not in emps:
            row_errors.append(f"الرقم {no} غير موجود أو غير نشط")
        elif no in seen:
            row_errors.append(f"الرقم {no} مكرّرٌ في الملفّ")
        seen.add(no)

        try:
            bal = Decimal(cell("balance").replace(",", ""))
            # ⚠️ **ورصيدٌ سالب يُرفض**: فمن استهلك أكثر ممّا استحقّ
            # **يُسوّى قبل النقل لا بعده**.
            if bal < 0:
                row_errors.append("الرصيد سالب — سوِّه قبل النقل")
        except (InvalidOperation, ValueError):
            bal = None
            row_errors.append(f"رصيد غير صالح: {cell('balance')}")

        as_of = None
        v = cell("as_of")
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                as_of = datetime.strptime(v, fmt).date()
                break
            except ValueError:
                continue
        if as_of is None:
            row_errors.append(f"تاريخ غير مفهوم: {v}")

        if row_errors:
            errors.append({"row": n, "employee_no": no or "—",
                           "errors": row_errors})
        else:
            parsed.append({"_row": n, "employment": emps[no],
                           "employee_no": no, "balance": bal,
                           "as_of": as_of})

    return {
        "rows": parsed, "errors": errors,
        "total": len(parsed) + len(errors),
        "valid": len(parsed), "invalid": len(errors),
        # ⚠️ **ولا استيراد بخطأ واحد** (قرار جواد)
        "can_import": not errors and bool(parsed),
    }


@transaction.atomic
def execute(*, company, leave_type, parsed_rows):
    """
    يكتب الأرصدة — **في معاملةٍ واحدة**.

    ⚠️⚠️ **والرصيد افتتاحيٌّ لا نهائيّ**: فالاستحقاق اليوميّ
    **يُضاف إليه** من تاريخه — **ولا يحلّ محلّه**.
    """
    from apps.leaves.models import LeaveBalance

    if not parsed_rows:
        raise BalanceImportError("لا صفوف للاستيراد")

    done = []
    for rec in parsed_rows:
        emp = rec["employment"]
        as_of = rec["as_of"]
        bal, _created = LeaveBalance.objects.update_or_create(
            employment=emp, leave_type=leave_type, year=as_of.year,
            defaults={
                "account_id": emp.account_id,
                "company_id": emp.company_id,
                "opening_balance": rec["balance"],
                # ⚠️ **والاحتساب يبدأ من تاريخه** — فلا يُعاد
                # استحقاقُ ما قبله
                "last_accrual_date": as_of,
            })
        done.append(rec["employee_no"])

    return {"imported": len(done), "employee_nos": done}
