"""
استيراد الحضور والانصراف (ق-171).

**بلاغ جواد:** شركةٌ تنتقل من نظامٍ آخر **لا تُعيد إدخال سجلّ
حضور موظفيها يومًا يومًا**.

⚠️ **والخطأ في سطرٍ يوقف الملفّ كلّه** (قرار جواد): فاستيرادٌ
نصفُه ناقصٌ **أسوأ من لا شيء**.

⚠️⚠️ **ولا يُستورَد يومٌ في مسيرٍ معتمد**: فالأجر احتُسب عليه،
**وتغييرُه يجعل القسيمة لا تفسّر نفسها**.
"""
import csv
import io
from datetime import date, datetime, time

from django.db import transaction


class DayImportError(Exception):
    """سببٌ يُعرض للمستخدم كما هو."""


COLUMNS = [
    ("employee_no", "الرقم الوظيفي", True, "مثال: 1001"),
    ("work_date", "التاريخ", True, "2025-12-01"),
    ("status", "الحالة", True, "present / absent / leave"),
    ("check_in", "الدخول", False, "08:00"),
    ("check_out", "الخروج", False, "17:00"),
]

#: ⚠️ **حالاتٌ معلَنة**: فحالةٌ مخترَعة **تُفسد التقارير صامتة**
STATUS_MAP = {
    "present": "present", "حاضر": "present",
    "absent": "absent", "غائب": "absent",
    "leave": "leave", "إجازة": "leave",
    "holiday": "holiday", "عطلة": "holiday",
    "weekend": "weekend", "راحة": "weekend",
}


def template_xlsx():
    """قالبٌ بعناوينه ومثالٍ يُحذف."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "الحضور"
    ws.sheet_view.rightToLeft = True

    head = PatternFill("solid", start_color="1F4E5F")
    req = PatternFill("solid", start_color="C0392B")

    for i, (key, label, required, example) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = req if required else head
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(
            len(label) + 8, 18)
        ex = ws.cell(row=2, column=i, value=example)
        ex.font = Font(italic=True, color="888888", size=10)
        ex.alignment = Alignment(horizontal="center")

    # ⚠️ **وقائمةٌ منسدلة للحالة** — فالعميل قد لا يعرف الصيغة
    dv = DataValidation(
        type="list",
        formula1='"present,absent,leave,holiday,weekend"',
        allow_blank=True, showDropDown=False)
    dv.error = "اختر من القائمة"
    ws.add_data_validation(dv)
    dv.add("C3:C5000")

    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_to_csv(data):
    """يحوّل xlsx لنصّ — ⚠️ **والأوقات تعود كائناتٍ**."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    buf = io.StringIO()
    w = csv.writer(buf)
    for row in ws.iter_rows(values_only=True):
        cells = []
        for v in row:
            if v is None:
                cells.append("")
            elif isinstance(v, datetime):
                cells.append(v.date().isoformat())
            elif isinstance(v, date):
                cells.append(v.isoformat())
            elif isinstance(v, time):
                cells.append(v.strftime("%H:%M"))
            elif isinstance(v, float) and v.is_integer():
                cells.append(str(int(v)))
            else:
                cells.append(str(v).strip())
        w.writerow(cells)
    return buf.getvalue()


def _parse_time(v):
    v = (v or "").strip()
    if not v:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(v, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"وقت غير مفهوم: {v}")


def parse_file(content, company):
    """
    يقرأ الملفّ ويفحصه سطرًا سطرًا — **بلا كتابة**.

    ⚠️ **ويُرجع كل الأخطاء لا أوّلها**.
    """
    from apps.employees.models import Employment, EmploymentStatus
    from apps.payroll.models import PayrollRun, PayrollRunStatus

    if isinstance(content, bytes) and content[:2] == b"PK":
        content = _xlsx_to_csv(content)
    if isinstance(content, bytes):
        for enc in ("utf-8-sig", "utf-8", "cp1256"):
            try:
                content = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise DayImportError(
                "تعذّر قراءة الملفّ — احفظه بترميز UTF-8")

    first = content.split("\n", 1)[0]
    delim = ";" if first.count(";") > first.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(content), delimiter=delim)
            if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise DayImportError("الملفّ فارغ")

    headers = [h.strip() for h in rows[0]]
    by_label = {c[1]: c[0] for c in COLUMNS}
    mapping = {by_label[h]: i for i, h in enumerate(headers)
               if h in by_label}
    missing = [c[1] for c in COLUMNS if c[2] and c[0] not in mapping]
    if missing:
        raise DayImportError(
            "أعمدة إلزامية ناقصة: " + "، ".join(missing))

    emps = {e.employee_no: e for e in Employment.objects.filter(
        company=company, status=EmploymentStatus.ACTIVE)}

    # ⚠️⚠️ **وشهورُ المسيرات المعتمدة محجوبة**
    locked = {(r.period_year, r.period_month) for r in
              PayrollRun.objects.filter(
                  company=company,
                  status__in=[PayrollRunStatus.APPROVED,
                              PayrollRunStatus.PAID])}

    parsed, errors, seen = [], [], set()
    for n, raw in enumerate(rows[1:], start=2):
        def cell(k):
            i = mapping.get(k)
            return (raw[i].strip() if i is not None and i < len(raw)
                    else "")

        no = cell("employee_no")
        if no.startswith("مثال"):
            continue

        row_errors = []
        if not no:
            row_errors.append("«الرقم الوظيفي» مطلوب")
        elif no not in emps:
            row_errors.append(f"الرقم {no} غير موجود أو غير نشط")

        work_date = None
        v = cell("work_date")
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                work_date = datetime.strptime(v, fmt).date()
                break
            except ValueError:
                continue
        if work_date is None:
            row_errors.append(f"تاريخ غير مفهوم: {v}")
        else:
            key = (no, work_date)
            if key in seen:
                row_errors.append(f"{no} مكرّرٌ في {work_date}")
            seen.add(key)
            if (work_date.year, work_date.month) in locked:
                row_errors.append(
                    f"شهر {work_date:%Y-%m} في مسيرٍ معتمد — "
                    "لا يُستورَد")

        st = STATUS_MAP.get(cell("status").strip().lower())
        if st is None:
            row_errors.append(f"حالة غير معروفة: {cell('status')}")

        cin = cout = None
        try:
            cin = _parse_time(cell("check_in"))
            cout = _parse_time(cell("check_out"))
        except ValueError as e:
            row_errors.append(str(e))

        # ⚠️ **وخروجٌ قبل دخول خطأ** — فاليوم لا يُقرأ عكسًا
        if cin and cout and cout <= cin:
            row_errors.append("الخروج قبل الدخول أو مساوٍ له")

        if row_errors:
            errors.append({"row": n, "employee_no": no or "—",
                           "errors": row_errors})
        else:
            parsed.append({"_row": n, "employment": emps[no],
                           "employee_no": no, "work_date": work_date,
                           "status": st, "check_in": cin,
                           "check_out": cout})

    return {
        "rows": parsed, "errors": errors,
        "total": len(parsed) + len(errors),
        "valid": len(parsed), "invalid": len(errors),
        "can_import": not errors and bool(parsed),
    }


@transaction.atomic
def execute(*, company, parsed_rows):
    """
    يكتب الأيام — **في معاملةٍ واحدة**.

    ⚠️ **وإخفاق سطرٍ يُرجع الجميع**.
    """
    from apps.attendance.models import AttendanceDay

    if not parsed_rows:
        raise DayImportError("لا صفوف للاستيراد")

    from django.utils import timezone

    def _stamp(day, t):
        """
        ⚠️ **والحقل تاريخٌ ووقت لا وقتًا**: فيُبنى من اليوم
        والساعة معًا — **وبلا هذا يُرفض الكتابة**.
        """
        if t is None:
            return None
        naive = datetime.combine(day, t)
        tz = timezone.get_current_timezone()
        return timezone.make_aware(naive, tz)

    n = 0
    for rec in parsed_rows:
        emp = rec["employment"]
        day = rec["work_date"]
        AttendanceDay.objects.update_or_create(
            employment=emp, work_date=day,
            defaults={
                "account_id": emp.account_id,
                "company_id": emp.company_id,
                "status": rec["status"],
                "first_in": _stamp(day, rec["check_in"]),
                "last_out": _stamp(day, rec["check_out"]),
            })
        n += 1
    return {"imported": n}
