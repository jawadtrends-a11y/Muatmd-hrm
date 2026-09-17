"""
الاستيراد الجماعيّ للموظفين (ق-154).

**شركةٌ بمئة موظف، أو منتقلةٌ من نظامٍ آخر** — لا يُعقل إدخالهم
واحدًا واحدًا (قرار جواد).

⚠️⚠️ **والخطأ في سطرٍ يوقف الملفّ كلّه**: فاستيرادٌ نصفُه ناقصٌ
**أسوأ من لا شيء** — والشركة لا تعرف من دخل ومن بقي.

⚠️ **والمعاينة قبل التنفيذ دائمًا**: فمن يرى ما سيقع يُصلحه قبل
أن يقع.
"""
import csv
import io
import logging
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

logger = logging.getLogger(__name__)


class ImportError_(Exception):
    """سببٌ يُعرض للمستخدم كما هو."""


#: أعمدة القالب — **الإلزاميّ أوّلًا**
#
# ⚠️ والعنوان عربيٌّ لأن من يملأ الملفّ موظفُ موارد لا مبرمج.
COLUMNS = [
    ("employee_no", "الرقم الوظيفي", True,
     "مثال: 1001 — ولا يتكرّر"),
    ("first_name_ar", "الاسم الأول", True, "محمد"),
    ("family_name_ar", "اسم العائلة", True, "العتيبي"),
    ("gender", "الجنس", True, "male أو female"),
    ("nationality_code", "الجنسية", True, "SA · EG · PK"),
    ("id_type", "نوع الهوية", True,
     "national_id للسعوديّ · iqama للمقيم"),
    ("id_number", "رقم الهوية", True, "1012345678"),
    ("join_date", "تاريخ المباشرة", True, "2026-01-15"),
    ("mobile", "الجوال", False, "0501234567"),
    ("email", "البريد", False, "name@company.com"),
    ("department", "الإدارة", False, "المالية — باسمها"),
    ("job_title", "المسمى الوظيفي", False, "محاسب"),
    ("basic_salary", "الراتب الأساسي", False, "8000"),
    ("housing_allowance", "بدل السكن", False, "2000"),
    ("transport_allowance", "بدل النقل", False, "800"),
    ("iban", "الآيبان", False, "SA0380000000608010167519"),

    # ق-216: ── حقولٌ كان الملفّ يفتقدها (بلاغ جواد) ──
    #
    # ⚠️⚠️ **فمن استورد ثلاثين موظفًا يُكملهم واحدًا واحدًا** —
    # والاستيراد يفقد معناه.
    #
    # ⚠️ **وما تُركَ عمدًا**: الهجريّ (**يُحتسب آليًّا**)، وحالةُ
    # الخدمة وانتهاؤها (**أحداثٌ لا بياناتُ تعيين**)، وتواريخُ
    # إصدار الهوية والجواز (**فالنظام يستعمل الانتهاء وحده**)،
    # والإجماليُّ والصافي والتأمينات (**محسوبةٌ، وإدخالها يخلق
    # تناقضًا**).

    # ── الأسماء الكاملة ──
    ("father_name_ar", "اسم الأب", False, "عبدالله"),
    ("grandfather_name_ar", "اسم الجد", False, "سعد"),
    ("full_name_en", "الاسم بالإنجليزية", False, "Mohammed Alotaibi"),

    # ── بيانات شخصية ──
    ("birth_date", "تاريخ الميلاد", False, "1990-05-20"),
    ("marital_status", "الحالة الاجتماعية", False,
     "single · married · divorced · widowed"),

    # ── الهوية والجواز ──
    ("id_expiry_date", "انتهاء الهوية", False, "2028-06-30"),
    ("passport_number", "رقم الجواز", False, "A12345678"),
    ("passport_expiry_date", "انتهاء الجواز", False, "2030-01-15"),

    # ── التنظيم ──
    #
    # ⚠️ **والمدير بالرقم الوظيفيّ لا بالاسم** (قرار جواد):
    # **فالأسماء تتكرّر والأرقام لا**.
    ("direct_manager_no", "الرقم الوظيفي للمدير المباشر", False,
     "1011 — في الملفّ أو في النظام"),
    ("branch", "الفرع", False, "الفرع الرئيسي — باسمه"),
    ("primary_site", "موقع العمل", False, "المقرّ — باسمه"),
    ("cost_center", "مركز التكلفة", False, "اختياريّ"),
    ("job_grade", "الدرجة الوظيفية", False, "اختياريّة"),
    ("job_step", "المرتبة الوظيفية", False, "اختياريّة"),

    # ── العقد والدوام ──
    ("employment_type", "نوع التوظيف", False,
     "full_time كامل · part_time جزئيّ"),
    ("contract_type", "نوع العقد", False,
     "fixed_term محدّد · unlimited غير محدّد"),
    ("contract_end_date", "انتهاء العقد", False,
     "للمحدّد المدّة فقط"),
    ("probation_days", "أيام التجربة", False, "90"),
    ("service_start_date", "بداية الخدمة المحتسبة", False,
     "اتركه فارغًا ليساوي تاريخ المباشرة"),

    # ── البصمة ──
    ("allow_mobile_punch", "بصمة الجوال", False, "نعم · لا"),

    # ── التأمينات وقوى ──
    ("is_gosi_registered", "مسجَّل بالتأمينات", False, "نعم · لا"),
    ("gosi_declared_wage", "الأجر المسجَّل بالتأمينات", False,
     "اتركه فارغًا ليساوي الأساسيّ + السكن"),
    ("gosi_establishment_no", "رقم منشأة التأمينات", False, ""),
    ("gosi_borne_by_company", "الشركة تتحمّل حصّة الموظف", False,
     "نعم · لا"),
    ("is_mol_registered", "مسجَّل في قوى", False, "نعم · لا"),
    ("mol_contract_no", "رقم عقد قوى", False, ""),

    # ── الصرف ──
    ("bank_code", "رمز البنك", False,
     "اتركه فارغًا ليُستخرج من الآيبان"),
    ("payment_method", "طريقة الصرف", False,
     "bank بنكيّ · cash نقدًا"),
    ("include_in_wps", "يدخل حماية الأجور", False, "نعم · لا"),
]

#: قوائمُ منسدلة في القالب — ⚠️ **فالعميل قد لا يعرف الصيغة**
#
# وبلا قائمةٍ يكتب «ذكر» فيُرفض الملفّ، ويُعيد الرفع مرارًا
# (بلاغ جواد).
DROPDOWNS = {
    "gender": ["male", "female"],
    "id_type": ["national_id", "iqama"],
}

#: ⚠️ **والجنسيات الشائعة لا كلّها**: فقائمةٌ بمئتَي دولة تُرهق،
# والحقل يبقى حرًّا لمن لم يجد جنسيّته.
NATIONALITIES = [
    ("SA", "السعودية"), ("EG", "مصر"), ("PK", "باكستان"),
    ("IN", "الهند"), ("BD", "بنغلاديش"), ("YE", "اليمن"),
    ("SD", "السودان"), ("SY", "سوريا"), ("JO", "الأردن"),
    ("PH", "الفلبين"), ("LK", "سريلانكا"), ("NP", "نيبال"),
    ("ID", "إندونيسيا"), ("LB", "لبنان"), ("PS", "فلسطين"),
    ("IQ", "العراق"), ("TN", "تونس"), ("MA", "المغرب"),
    ("DZ", "الجزائر"), ("ET", "إثيوبيا"), ("KE", "كينيا"),
    ("NG", "نيجيريا"), ("TR", "تركيا"), ("GB", "بريطانيا"),
    ("US", "أمريكا"),
]


#: بنود الأجر التي تُستورَد — **ورمزُها في النظام**
SALARY_MAP = {
    "basic_salary": "BASIC",
    "housing_allowance": "HOUSING",
    "transport_allowance": "TRANSPORT",
}


def _ref_code(prefix, name, existing):
    """
    رمزٌ لمرجعٍ يُنشأ ضمنًا (ق-215).

    ⚠️ **ولا يتكرّر**: فرمزان متطابقان **يكسران قيد التفرّد** —
    والعدّاد يمضي حتى يجد فراغًا.
    """
    import re as _re

    base = _re.sub(r"[^A-Za-z0-9]+", "", (name or "").upper())[:6]
    if not base:
        base = prefix
    taken = {getattr(v, "code", "") for v in existing.values()}
    code, i = base, 1
    while code in taken:
        i += 1
        code = f"{base}{i}"
    return code[:20]


def template_xlsx():
    """
    قالبٌ بصيغة Excel — **لا CSV**.

    ⚠️⚠️ **فإكسل العربيّ يتوقّع الفاصلة المنقوطة**: ويفتح ملفّ
    CSV بالفواصل **في عمودٍ واحد**، فيظنّ العميل القالب معطوبًا.

    **والـxlsx لا فواصل فيه ولا ترميز** — يفتح كما هو في كل بيئة.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "الموظفون"
    ws.sheet_view.rightToLeft = True

    head_fill = PatternFill("solid", start_color="1F4E5F")
    req_fill = PatternFill("solid", start_color="C0392B")
    bold = Font(bold=True, color="FFFFFF", size=11)

    for i, (key, label, required, example) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = bold
        # ⚠️ **والإلزاميّ بلونٍ مختلف** — فيُرى قبل أن يُقرأ
        c.fill = req_fill if required else head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = max(
            len(label) + 6, 16)

        ex = ws.cell(row=2, column=i, value=example)
        ex.font = Font(italic=True, color="888888", size=10)
        ex.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 26

    # ── القوائم المنسدلة (ق-155) ──
    #
    # ⚠️ **فالعميل قد لا يعرف الصيغة**: يكتب «ذكر» فيُرفض الملفّ،
    # ويُعيد الرفع مرارًا (بلاغ جواد).
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    keys = [c[0] for c in COLUMNS]

    for key, options in DROPDOWNS.items():
        if key not in keys:
            continue
        col = get_column_letter(keys.index(key) + 1)
        dv = DataValidation(
            type="list", formula1='"' + ",".join(options) + '"',
            allow_blank=True, showDropDown=False)
        dv.error = "اختر من القائمة"
        dv.errorTitle = "قيمة غير مقبولة"
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}1000")

    # ⚠️ **والجنسيات في ورقةٍ مرجعية**: فقائمةٌ طويلة لا تُكتب في
    # الصيغة، والورقة تُخفى فلا تُربك من يملأ.
    ref = wb.create_sheet("مرجع")
    ref["A1"] = "الرمز"
    ref["B1"] = "الجنسية"
    for i, (code, name) in enumerate(NATIONALITIES, 2):
        ref[f"A{i}"] = code
        ref[f"B{i}"] = name
    ref.sheet_state = "hidden"

    if "nationality_code" in keys:
        col = get_column_letter(keys.index("nationality_code") + 1)
        dv = DataValidation(
            type="list",
            formula1=f"مرجع!$A$2:$A${len(NATIONALITIES) + 1}",
            allow_blank=True, showDropDown=False)
        # ⚠️ **ولا يُمنع غيرُها**: فمن لم يجد جنسيّته يكتبها
        dv.showErrorMessage = False
        dv.promptTitle = "الجنسية"
        dv.prompt = "اختر من القائمة أو اكتب رمز الدولة بحرفين"
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}1000")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_csv():
    """CSV احتياطيّ — ⚠️ وبالفاصلة المنقوطة لإكسل العربيّ."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([c[1] for c in COLUMNS])
    w.writerow([c[3] for c in COLUMNS])
    return buf.getvalue()


def _xlsx_to_csv(data):
    """
    يحوّل xlsx إلى نصٍّ يُقرأ بالمنطق نفسه.

    ⚠️ **فالتواريخ تعود كائناتٍ لا نصًّا** — وتُنسَّق هنا.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active

    buf = io.StringIO()
    w = csv.writer(buf)
    for row in ws.iter_rows(values_only=True):
        cells = []
        for v in row:
            if v is None:
                cells.append("")
            elif isinstance(v, datetime):
                cells.append(v.date().isoformat())
            elif isinstance(v, date):
                cells.append(v.isoformat())
            elif isinstance(v, float) and v.is_integer():
                cells.append(str(int(v)))
            else:
                cells.append(str(v).strip())
        w.writerow(cells)
    return buf.getvalue()


def _parse_date(v):
    v = (v or "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"تاريخ غير مفهوم: {v}")


def _parse_amount(v):
    v = (v or "").strip().replace(",", "")
    if not v:
        return None
    try:
        return Decimal(v)
    except InvalidOperation:
        raise ValueError(f"مبلغ غير صالح: {v}")


def parse_file(content, company):
    """
    يقرأ الملفّ ويفحصه **سطرًا سطرًا** — بلا إنشاء.

    ⚠️ **ويُرجع كل الأخطاء لا أوّلها**: فمن يُصلح خطأً ثم يرفع
    ليكتشف ثانيًا **يكرّر الرفع عشرًا**.
    """
    from apps.employees.models import Employment, Person
    from apps.organization.models import Department, JobTitle

    # ⚠️ **والقراءة تقبل xlsx وCSV**: فالعميل قد يحفظ بأيّهما
    if isinstance(content, bytes) and content[:2] == b"PK":
        content = _xlsx_to_csv(content)

    if isinstance(content, bytes):
        for enc in ("utf-8-sig", "utf-8", "cp1256"):
            try:
                content = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportError_("تعذّر قراءة الملفّ — احفظه بترميز UTF-8")

    # ⚠️ **والفاصل يُكتشف**: فمن حفظ من إكسل العربيّ يحصل على «؛»
    first = content.split("\n", 1)[0]
    delim = ";" if first.count(";") > first.count(",") else ","
    reader = csv.reader(io.StringIO(content), delimiter=delim)
    rows = [r for r in reader if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise ImportError_("الملفّ فارغ — عبّئ صفًّا واحدًا على الأقلّ")

    headers = [h.strip() for h in rows[0]]
    by_label = {c[1]: c[0] for c in COLUMNS}
    mapping = {}
    for i, h in enumerate(headers):
        if h in by_label:
            mapping[by_label[h]] = i

    missing_cols = [c[1] for c in COLUMNS
                    if c[2] and c[0] not in mapping]
    if missing_cols:
        raise ImportError_(
            "أعمدة إلزامية ناقصة: " + "، ".join(missing_cols))

    # مراجع الشركة — تُقرأ مرّةً لا لكل سطر
    depts = {d.name_ar.strip(): d for d in
             Department.objects.filter(company=company)}
    titles = {t.name_ar.strip(): t for t in
              JobTitle.objects.filter(company=company)}
    existing_nos = set(Employment.objects.filter(
        company=company).values_list("employee_no", flat=True))
    existing_ids = set(Person.objects.filter(
        account=company.account).values_list("id_number", flat=True))

    parsed, errors = [], []
    seen_nos, seen_ids = set(), set()

    # ق-215: **ما سيُنشأ** — يُعرَض في المعاينة لا يُنشأ فيها
    will_create_depts, will_create_titles = set(), set()

    for n, raw in enumerate(rows[1:], start=2):
        def cell(key):
            i = mapping.get(key)
            return (raw[i].strip() if i is not None and i < len(raw)
                    else "")

        # ⚠️⚠️ **وسطر المثال يُعلَّم بعلامةٍ لا تلتبس ببيانات**:
        # فمطابقةُ القيم **تحذف صفًّا حقيقيًّا صامتًا**.
        if cell("employee_no").startswith("مثال"):
            continue

        rec, row_errors = {"_row": n}, []

        for key, label, required, _ in COLUMNS:
            v = cell(key)
            if required and not v:
                row_errors.append(f"«{label}» مطلوب")
            rec[key] = v

        # ── تكرارٌ داخل الملفّ أو في النظام ──
        no = rec.get("employee_no")
        if no:
            if no in seen_nos:
                row_errors.append(f"الرقم الوظيفي {no} مكرّرٌ في الملفّ")
            elif no in existing_nos:
                row_errors.append(f"الرقم الوظيفي {no} مستعملٌ في النظام")
            seen_nos.add(no)

        idn = rec.get("id_number")
        if idn:
            if idn in seen_ids:
                row_errors.append("رقم الهوية مكرّرٌ في الملفّ")
            elif idn in existing_ids:
                row_errors.append("رقم الهوية مسجَّلٌ في النظام")
            seen_ids.add(idn)

        if rec.get("gender") not in ("male", "female", ""):
            row_errors.append("الجنس: male أو female")
        if rec.get("id_type") not in ("national_id", "iqama", ""):
            row_errors.append("نوع الهوية: national_id أو iqama")

        try:
            rec["join_date"] = _parse_date(rec.get("join_date"))

            # ق-216: ⚠️ **وبقيّة التواريخ تُحوَّل كذلك**: فالمعاينة
            # تقرأ كل عمودٍ **نصًّا**، والنماذج تريد تاريخًا —
            # **ونصٌّ في موضع تاريخ يُسقط الاستيراد**.
            for _dk in ("birth_date", "id_expiry_date",
                        "passport_expiry_date", "contract_end_date",
                        "service_start_date"):
                _raw = rec.get(_dk)
                if _raw:
                    rec[_dk] = _parse_date(_raw)
        except ValueError as e:
            row_errors.append(str(e))

        for k in SALARY_MAP:
            try:
                rec[k] = _parse_amount(rec.get(k))
            except ValueError as e:
                row_errors.append(str(e))

        # ق-215: ── المراجع: **تُطابق بالاسم، وما لا يوجد يُنشأ** ──
        #
        # **قرار جواد:** «المفترض يتم اختراع وقت الإنشاء — بدون
        # تكرار».
        #
        # ⚠️⚠️ **والخطأ الإملائيّ يصير إدارةً جديدة** — فالمعاينة
        # **تُنبّه بما سيُنشأ قبل أن يُنشأ**، والعميل يراجع.
        d = rec.get("department")
        if d and d not in depts:
            will_create_depts.add(d)
        t = rec.get("job_title")
        if t and t not in titles:
            will_create_titles.add(t)

        if row_errors:
            errors.append({"row": n,
                           "employee_no": no or "—",
                           "errors": row_errors})
        else:
            parsed.append(rec)

    return {
        "rows": parsed,
        "errors": errors,
        "total": len(parsed) + len(errors),
        "valid": len(parsed),
        "invalid": len(errors),
        # ⚠️ **ولا استيراد بخطأ واحد** (قرار جواد)
        "can_import": not errors and bool(parsed),

        # ق-215: ⚠️⚠️ **وما سيُنشأ يُعرَض قبل أن يُنشأ**: فالخطأ
        # الإملائيّ **يصير إدارةً جديدة** — والعميل يراجع القائمة
        # قبل التنفيذ.
        "new_departments": sorted(will_create_depts),
        "new_job_titles": sorted(will_create_titles),
    }


#: حقولُ الشخص الاختيارية — تُمرَّر إن وُجدت
_PERSON_KEYS = (
    "father_name_ar", "grandfather_name_ar", "full_name_en",
    "birth_date", "marital_status", "id_expiry_date",
    "passport_number", "passport_expiry_date",
)

#: حقولُ الارتباط النصّية والتاريخية
_EMP_KEYS = (
    "employment_type", "contract_type", "contract_end_date",
    "service_start_date", "gosi_establishment_no", "mol_contract_no",
    "payment_method", "bank_code",
)

#: حقولٌ نعم/لا — ⚠️ **والعميل يكتب «نعم» لا `true`**
_EMP_BOOLS = (
    "allow_mobile_punch", "is_gosi_registered", "is_mol_registered",
    "include_in_wps", "gosi_borne_by_company",
)


def _yes(v):
    """«نعم» و«true» و«1» — ⚠️ **والفراغ ليس «لا» بل «لم يُذكر»**."""
    s = (v or "").strip().lower()
    if not s:
        return None
    return s in ("نعم", "yes", "true", "1", "y", "✓")


def _person_extra(rec):
    """ما يُمرَّر لـ`create_person` من الحقول الاختيارية."""
    out = {}
    for k in _PERSON_KEYS:
        v = rec.get(k)
        if v not in (None, ""):
            out[k] = v
    return out


@transaction.atomic
def execute(*, company, parsed_rows, by_person_id=None):
    """
    يُنشئ الموظفين — **في معاملةٍ واحدة**.

    ⚠️ **فإخفاق سطرٍ يُرجع الجميع**: وإلا بقي نصفُ الملفّ في
    النظام ولا يُعرف أين توقّف.
    """
    from apps.employees.services.hiring import (
        create_employment, create_person)
    from apps.organization.models import Department, JobTitle
    from apps.payroll.models import PayComponent

    if not parsed_rows:
        raise ImportError_("لا صفوف للاستيراد")

    comps = {c.code: c for c in PayComponent.objects.filter(
        company=company)}
    depts = {d.name_ar.strip(): d for d in
             Department.objects.filter(company=company)}
    titles = {t.name_ar.strip(): t for t in
              JobTitle.objects.filter(company=company)}

    from apps.organization.models import Branch, CostCenter

    branches = {b.name_ar.strip(): b for b in
                Branch.objects.filter(company=company)}
    centers = {c.name_ar.strip(): c for c in
               CostCenter.objects.filter(company=company)}

    created = []
    pending_managers = []
    for rec in parsed_rows:
        try:
            person, _ = create_person(
                account=company.account,
                first_name_ar=rec["first_name_ar"],
                family_name_ar=rec["family_name_ar"],
                gender=rec["gender"],
                nationality_code=rec["nationality_code"],
                id_type=rec["id_type"], id_number=rec["id_number"],
                mobile=rec.get("mobile", ""),
                email=rec.get("email", ""),
                # ق-216: **وبقيّة بيانات الشخص** — فالقالب
                # يحملها، **وإهمالُها يعد بما لا يقع**.
                **_person_extra(rec),
                # ⚠️ تحذير التشابه يُتجاوز في الاستيراد الجماعيّ:
                # فمئة موظفٍ فيهم متشابهو الأسماء بالضرورة.
                force=True)

            lines = []
            for key, code in SALARY_MAP.items():
                amt = rec.get(key)
                if amt and code in comps:
                    lines.append((comps[code], amt))

            extra = {}

            # ق-215: ⚠️⚠️ **وما لا يوجد يُنشأ** (قرار جواد) —
            # **بلا تكرار**: فالقاموس يُحدَّث فورًا، **فإدارةٌ
            # على عشرين صفًّا تُنشأ مرّةً واحدة**.
            d = rec.get("department")
            if d:
                if d not in depts:
                    depts[d] = Department.objects.create(
                        account=company.account, company=company,
                        name_ar=d, code=_ref_code("DEP", d, depts))
                extra["department"] = depts[d]

            t = rec.get("job_title")
            if t:
                if t not in titles:
                    # ⚠️ **والمسمّى بلا رمز** — اسمه يكفي
                    titles[t] = JobTitle.objects.create(
                        account=company.account, company=company,
                        name_ar=t)
                extra["job_title"] = titles[t]

            # ق-216: **بقيّة بيانات الارتباط**
            for k in _EMP_KEYS:
                v = rec.get(k)
                if v not in (None, ""):
                    extra[k] = v

            for k in _EMP_BOOLS:
                b = _yes(rec.get(k))
                if b is not None:
                    extra[k] = b

            # ⚠️⚠️ **والأرقام تُحوَّل**: فالمعاينة تقرأ كل عمودٍ
            # **نصًّا**، و`probation_days` نصًّا **يكسر حساب
            # timedelta**.
            v = rec.get("probation_days")
            if v not in (None, ""):
                try:
                    extra["probation_days"] = int(float(str(v).strip()))
                except (TypeError, ValueError):
                    pass

            v = rec.get("gosi_declared_wage")
            if v not in (None, ""):
                try:
                    extra["gosi_declared_wage"] = Decimal(
                        str(v).replace(",", "").strip())
                except (TypeError, ValueError, InvalidOperation):
                    pass

            # ── المراجع النصّية: تُنشأ إن لم توجد ──
            br = rec.get("branch")
            if br:
                if br not in branches:
                    branches[br] = Branch.objects.create(
                        account=company.account, company=company,
                        name_ar=br, code=_ref_code("BR", br, branches))
                extra["branch"] = branches[br]

            cc = rec.get("cost_center")
            if cc:
                if cc not in centers:
                    centers[cc] = CostCenter.objects.create(
                        account=company.account, company=company,
                        name_ar=cc, code=_ref_code("CC", cc, centers))
                extra["cost_center"] = centers[cc]

            emp, _, _ = create_employment(
                person=person, company=company,
                employee_no=rec["employee_no"],
                join_date=rec["join_date"],
                salary_lines=lines or None,
                iban=rec.get("iban", ""), **extra)
            created.append(emp.employee_no)

            # ⚠️⚠️ **والمدير يُربط بعد إنشاء الجميع**: فمديرٌ
            # **في الملفّ نفسه قد يأتي بعد مرؤوسه**.
            mgr = rec.get("direct_manager_no")
            if mgr:
                pending_managers.append((emp, mgr))
        except Exception as e:          # noqa: BLE001
            # ⚠️ **والمعاملة تُرجع الجميع** — فالرسالة تُسمّي السطر
            raise ImportError_(
                f"توقّف الاستيراد عند السطر {rec['_row']} "
                f"({rec.get('employee_no')}): {e}")

    # ق-216: ⚠️⚠️ **والمدير يُربط بعد إنشاء الجميع** (قرار جواد:
    # بالرقم الوظيفيّ): **فمديرٌ في الملفّ نفسه قد يأتي بعد
    # مرؤوسه** — والربط في حينه يُخفق.
    #
    # ⚠️ **ورقمٌ لا وجود له يُتجاوز**: فالموظف أُنشئ، **وربطُ
    # مديرٍ مفقودٍ لا يستحقّ إسقاط الاستيراد كلّه**.
    if pending_managers:
        from apps.employees.models import Employment as _Emp

        by_no = {e.employee_no: e
                 for e in _Emp.objects.filter(company=company)}
        for emp, mgr_no in pending_managers:
            mgr = by_no.get(str(mgr_no).strip())
            if mgr is not None and mgr.id != emp.id:
                emp.direct_manager = mgr
                emp.save(update_fields=["direct_manager", "updated_at"])

    logger.info("استُورد %s موظفًا للشركة %s", len(created), company.id)
    return {"created": len(created), "employee_nos": created}
